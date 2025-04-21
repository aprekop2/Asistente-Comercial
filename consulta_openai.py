### consulta_openai.py
import os
from openai import OpenAI

client = OpenAI(api_key=os.getenv("OPENAI_API_KEY"))

def consulta_chatgpt(pregunta):
    respuesta = client.chat.completions.create(
        model="gpt-3.5-turbo",
        messages=[
            {"role": "system", "content": "Eres un asistente comercial muy útil."},
            {"role": "user", "content": pregunta}
        ]
    )
    return respuesta.choices[0].message.content


### generador_docs.py
from docx import Document
import openpyxl

def generar_docx(nombre_archivo, contenido):
    doc = Document()
    doc.add_paragraph(contenido)
    doc.save(nombre_archivo)

def generar_excel(nombre_archivo, encabezados, datos):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(encabezados)
    for fila in datos:
        ws.append(fila)
    wb.save(nombre_archivo)


### supabase_config.py
import os
from supabase import create_client, Client
from datetime import datetime

SUPABASE_URL = os.getenv("SUPABASE_URL")
SUPABASE_KEY = os.getenv("SUPABASE_KEY")

supabase: Client = create_client(SUPABASE_URL, SUPABASE_KEY)

def guardar_interaccion(pregunta, respuesta):
    data = {
        "pregunta": pregunta,
        "respuesta": respuesta,
        "fecha": datetime.utcnow().isoformat()
    }
    supabase.table("interacciones").insert(data).execute()


### templates/index.html
