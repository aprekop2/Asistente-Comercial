from docx import Document
from openpyxl import Workbook

def generar_docx(texto, nombre_archivo="salida.docx"):
    doc = Document()
    doc.add_paragraph(texto)
    doc.save(nombre_archivo)
    return nombre_archivo

def generar_excel(data, nombre_archivo="salida.xlsx"):
    wb = Workbook()
    ws = wb.active
    for fila in data:
        ws.append(fila)
    wb.save(nombre_archivo)
    return nombre_archivo
